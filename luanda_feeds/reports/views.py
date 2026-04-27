"""
Reports app views for sales reporting and exports.
"""

from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import HttpResponse
from django.db.models import Sum, F
from datetime import datetime
from decimal import Decimal
from io import BytesIO

from sales.models import Sale, SaleItem

@login_required
def sales_report(request):
    """Generate sales report with date range filter."""
    if not request.user.can_view_reports():
        messages.error(request, 'You do not have permission to view reports.')
        return render(request, 'sales/cashier_dashboard.html')

    # Get date range from GET params
    start_date = request.GET.get('start_date')
    end_date = request.GET.get('end_date')

    if start_date and end_date:
        try:
            start_dt = datetime.strptime(start_date, '%Y-%m-%d')
            end_dt = datetime.strptime(end_date, '%Y-%m-%d')
            # Include end date full day
            end_dt = end_dt.replace(hour=23, minute=59, second=59)
        except ValueError:
            messages.error(request, 'Invalid date format.')
            start_date = end_date = None
    else:
        # Default to today
        now = datetime.now()
        start_dt = now.replace(hour=0, minute=0, second=0)
        end_dt = now.replace(hour=23, minute=59, second=59)
        start_date = now.strftime('%Y-%m-%d')
        end_date = now.strftime('%Y-%m-%d')

    # Query sales
    sales = Sale.objects.filter(timestamp__gte=start_dt, timestamp__lte=end_dt)

    # Aggregate by item
    item_sales = SaleItem.objects.filter(
        sale__timestamp__gte=start_dt,
        sale__timestamp__lte=end_dt
    ).values(
        'item_name'
    ).annotate(
        total_quantity=Sum('quantity'),
        total_revenue=Sum(F('quantity') * F('selling_price')),
        total_cost=Sum(F('quantity') * F('buying_price')),
    ).order_by('-total_quantity')

    # Calculate totals
    total_revenue = sum(item['total_revenue'] or Decimal('0') for item in item_sales)
    total_profit = sum(
        (item['total_revenue'] or Decimal('0')) - (item['total_cost'] or Decimal('0'))
        for item in item_sales
    )
    total_quantity = sum(item['total_quantity'] or Decimal('0') for item in item_sales)

    context = {
        'start_date': start_date,
        'end_date': end_date,
        'item_sales': item_sales,
        'total_revenue': total_revenue,
        'total_profit': total_profit,
        'total_quantity': total_quantity,
        'has_data': bool(item_sales),
    }

    # Store report data in session for export
    request.session['report_data'] = {
        'start_date': start_date,
        'end_date': end_date,
        'items': list(item_sales),
        'total_revenue': str(total_revenue),
        'total_profit': str(total_profit),
        'total_quantity': str(total_quantity),
    }

    return render(request, 'reports/sales_report.html', context)

@login_required
def export_pdf(request):
    """Export sales report to PDF."""
    if not request.user.can_view_reports():
        return HttpResponse('Permission denied', status=403)

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.enums import TA_CENTER

    # Get report data from session
    report_data = request.session.get('report_data', {})
    if not report_data:
        messages.error(request, 'No report data to export. Please generate a report first.')
        return HttpResponse('No report data', status=400)

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=0.5*inch, bottomMargin=0.5*inch)

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        alignment=TA_CENTER,
        fontSize=16,
    )
    subtitle_style = ParagraphStyle(
        'CustomSubtitle',
        parent=styles['Normal'],
        alignment=TA_CENTER,
        fontSize=10,
    )

    elements = []

    # Title
    elements.append(Paragraph("Luanda Feeds Centre", title_style))
    elements.append(Paragraph("Sales Report", title_style))
    elements.append(Spacer(1, 0.2*inch))
    elements.append(Paragraph(
        f"Date Range: {report_data['start_date']} to {report_data['end_date']}",
        subtitle_style
    ))
    elements.append(Spacer(1, 0.3*inch))

    # Table data
    table_data = [['Item Name', 'Quantity Sold', 'Revenue (KES)', 'Profit (KES)']]

    for item in report_data['items']:
        revenue = item['total_revenue'] or Decimal('0')
        cost = item['total_cost'] or Decimal('0')
        profit = revenue - cost

        table_data.append([
            item['item_name'][:30],
            f"{item['total_quantity']:.2f}",
            f"{revenue:,.2f}",
            f"{profit:,.2f}",
        ])

    # Add totals row
    table_data.append([
        'TOTAL',
        f"{Decimal(report_data['total_quantity']):.2f}",
        f"{Decimal(report_data['total_revenue']):,.2f}",
        f"{Decimal(report_data['total_profit']):,.2f}",
    ])

    # Create table
    table = Table(table_data, colWidths=[2.5*inch, 1.2*inch, 1.5*inch, 1.5*inch])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#00FFFF')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.black),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 1), (-1, -2), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 9),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('TOPPADDING', (0, 0), (-1, 0), 12),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#E0E0E0')),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
    ]))

    elements.append(table)
    elements.append(Spacer(1, 0.5*inch))

    # Summary
    summary_style = ParagraphStyle(
        'Summary',
        parent=styles['Normal'],
        fontSize=10,
    )
    elements.append(Paragraph(
        f"Total Revenue: KES {Decimal(report_data['total_revenue']):,.2f}",
        summary_style
    ))
    elements.append(Paragraph(
        f"Total Profit: KES {Decimal(report_data['total_profit']):,.2f}",
        summary_style
    ))

    doc.build(elements)

    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="sales_report_{report_data["start_date"]}_{report_data["end_date"]}.pdf"'
    return response

@login_required
def export_excel(request):
    """Export sales report to Excel."""
    if not request.user.can_view_reports():
        return HttpResponse('Permission denied', status=403)

    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    # Get report data from session
    report_data = request.session.get('report_data', {})
    if not report_data:
        messages.error(request, 'No report data to export. Please generate a report first.')
        return HttpResponse('No report data', status=400)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sales Report"

    # Styles
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="00FFFF", end_color="00FFFF", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Title
    ws.merge_cells('A1:D1')
    ws['A1'] = 'Luanda Feeds Centre - Sales Report'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = Alignment(horizontal="center")

    ws.merge_cells('A2:D2')
    ws['A2'] = f"Date Range: {report_data['start_date']} to {report_data['end_date']}"
    ws['A2'].alignment = Alignment(horizontal="center")

    # Headers
    headers = ['Item Name', 'Quantity Sold', 'Revenue (KES)', 'Profit (KES)']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # Data rows
    for row, item in enumerate(report_data['items'], 5):
        revenue = item['total_revenue'] or Decimal('0')
        cost = item['total_cost'] or Decimal('0')
        profit = revenue - cost

        ws.cell(row=row, column=1, value=item['item_name'][:30]).border = thin_border
        ws.cell(row=row, column=2, value=float(item['total_quantity'])).border = thin_border
        ws.cell(row=row, column=3, value=float(revenue)).border = thin_border
        ws.cell(row=row, column=4, value=float(profit)).border = thin_border

        # Number formatting
        ws.cell(row=row, column=2).number_format = '#,##0.00'
        ws.cell(row=row, column=3).number_format = '#,##0.00'
        ws.cell(row=row, column=4).number_format = '#,##0.00'

    # Totals row
    totals_row = len(report_data['items']) + 5
    ws.cell(row=totals_row, column=1, value="TOTAL").font = Font(bold=True)
    ws.cell(row=totals_row, column=2, value=float(Decimal(report_data['total_quantity']))).font = Font(bold=True)
    ws.cell(row=totals_row, column=3, value=float(Decimal(report_data['total_revenue']))).font = Font(bold=True)
    ws.cell(row=totals_row, column=4, value=float(Decimal(report_data['total_profit']))).font = Font(bold=True)

    for col in range(1, 5):
        ws.cell(row=totals_row, column=col).border = thin_border
        ws.cell(row=totals_row, column=col).fill = PatternFill(
            start_color="E0E0E0", end_color="E0E0E0", fill_type="solid"
        )
        if col > 1:
            ws.cell(row=totals_row, column=col).number_format = '#,##0.00'

    # Column widths
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 18

    from io import BytesIO
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="sales_report_{report_data["start_date"]}_{report_data["end_date"]}.xlsx"'
    return response
